"""Exercise every document parser, including the rejection paths.

    python scripts/smoke_parsers.py

Sample files are generated here rather than committed, so the test has no
fixtures to keep in sync and covers the awkward cases explicitly: a Word file
with real heading styles, a Word file with none, a PDF with a text layer, and a
PDF with no extractable text (the scanned-document case).
"""

from __future__ import annotations

import sys as _sys

for _stream in (_sys.stdout, _sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

import io
import sys

from app.ingest.parsers import (
    EmptyDocument,
    UnsupportedDocument,
    parse_to_markdown,
)

LOREM = (
    "Singapore's hawker centres are open-air complexes housing many stalls "
    "selling inexpensive local food. Chinatown Complex is the largest, with "
    "over two hundred stalls. Maxwell Food Centre is known for chicken rice. "
    "Lau Pa Sat occupies a restored Victorian cast-iron market building. "
)


def make_docx(with_headings: bool) -> bytes:
    from docx import Document

    document = Document()
    if with_headings:
        document.add_heading("Singapore Food Guide", level=0)
        document.add_heading("Hawker Centres", level=1)
        document.add_paragraph(LOREM)
        document.add_heading("Chinatown Complex", level=2)
        document.add_paragraph(LOREM)
        document.add_heading("Getting There", level=2)
        document.add_paragraph(LOREM)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Stall"
        table.cell(0, 1).text = "Dish"
        table.cell(1, 0).text = "Tian Tian"
        table.cell(1, 1).text = "Chicken rice"
    else:
        for _ in range(4):
            document.add_paragraph(LOREM)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(with_text: bool) -> bytes:
    """A minimal PDF, written by hand to avoid another dependency."""
    if with_text:
        lines = [
            "BT /F1 12 Tf 40 750 Td (Singapore Travel Notes) Tj ET",
        ]
        y = 720
        for sentence in LOREM.split(". "):
            if not sentence.strip():
                continue
            text = sentence.replace("(", "").replace(")", "").strip()
            lines.append(f"BT /F1 10 Tf 40 {y} Td ({text}.) Tj ET")
            y -= 18
        content = "\n".join(lines)
    else:
        # A page with graphics only -- the scanned-document case.
        content = "0.5 0.5 0.5 rg 40 400 500 300 re f"

    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n".encode()
    )
    return out.getvalue()


def main() -> int:
    failures: list[str] = []

    print("=" * 78)
    print("FORMATS THAT MUST PARSE")
    print("=" * 78)
    good_cases = [
        ("guide.md", f"# Singapore Guide\n\n## Hawkers\n\n{LOREM}\n\n### Maxwell\n\n{LOREM}".encode()),
        ("notes.txt", f"Singapore Travel Notes\n\n{LOREM}\n\n{LOREM}".encode()),
        ("page.html",
         f"<html><body><nav>skip me</nav><main><h1>Singapore</h1>"
         f"<h2>Food</h2><p>{LOREM}</p><h2>Transport</h2><p>{LOREM}</p>"
         f"</main><footer>skip</footer></body></html>".encode()),
        ("styled.docx", make_docx(with_headings=True)),
        ("unstyled.docx", make_docx(with_headings=False)),
        ("notes.pdf", make_pdf(with_text=True)),
    ]
    for filename, data in good_cases:
        try:
            parsed = parse_to_markdown(data, filename)
        except (EmptyDocument, UnsupportedDocument) as exc:
            print(f"\nFAIL {filename}: {exc}")
            failures.append(f"{filename} should have parsed: {exc}")
            continue
        print(f"\nOK   {filename}  [{parsed.format}]")
        print(f"       title    : {parsed.title}")
        print(f"       chars    : {parsed.char_count:,}")
        print(f"       headings : {parsed.headings[:4]}")
        for warning in parsed.warnings:
            print(f"       warning  : {warning[:100]}")
        if filename == "styled.docx":
            levels = [h for h in parsed.markdown.splitlines() if h.startswith("#")]
            # level=0 -> Title -> '#', level=1 -> '##', level=2 -> '###'
            depths = sorted({len(line.split(' ')[0]) for line in levels})
            has_nesting = len(depths) >= 3
            print('       heading depths present: ' + str(depths)
                  + ' (3+ levels: ' + str(has_nesting) + ')')
            if not has_nesting:
                failures.append('styled.docx lost its heading hierarchy')
            if "| Tian Tian | Chicken rice |" not in parsed.markdown:
                failures.append("styled.docx lost its table")
            else:
                print("       table flattened to pipe rows: True")
        if filename == "unstyled.docx" and not parsed.warnings:
            failures.append("unstyled.docx should warn about missing headings")
        if filename == "page.html":
            if "skip me" in parsed.markdown or "skip" in parsed.markdown.lower()[-30:]:
                failures.append("page.html kept nav/footer chrome")
            else:
                print("       nav/footer stripped: True")
        if not parsed.headings:
            failures.append(f"{filename} produced no headings")

    print()
    print("=" * 78)
    print("FILES THAT MUST BE REJECTED, NOT SILENTLY ACCEPTED")
    print("=" * 78)
    bad_cases = [
        ("scanned.pdf", make_pdf(with_text=False), EmptyDocument,
         "image-only PDF (the scanned-document case)"),
        ("tiny.md", b"# Hi", EmptyDocument, "too little text to answer anything"),
        ("archive.zip", b"PK\x03\x04 not really a zip", UnsupportedDocument,
         "unsupported extension"),
        ("noextension", b"x" * 500, UnsupportedDocument, "no extension at all"),
        ("broken.docx", b"this is not a docx at all", EmptyDocument,
         "corrupt file of a supported type"),
    ]
    for filename, data, expected, label in bad_cases:
        try:
            parse_to_markdown(data, filename)
        except expected as exc:
            print(f"\nOK   {filename} -- {label}")
            print(f"       {str(exc)[:150]}")
        except Exception as exc:
            print(f"\nFAIL {filename}: raised {type(exc).__name__}, "
                  f"expected {expected.__name__}")
            failures.append(f"{filename} raised the wrong error type")
        else:
            print(f"\nFAIL {filename} was accepted but should have been rejected")
            failures.append(f"{filename} was not rejected")

    print()
    print("=" * 78)
    if failures:
        print(f"{len(failures)} FAILURE(S):")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("All parser checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
