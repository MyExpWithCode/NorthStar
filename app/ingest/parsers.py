"""Convert an uploaded or fetched document into markdown.

One entry point, `parse_to_markdown`, dispatching on file extension. Every
parser has the same obligation: **preserve the heading hierarchy**, because the
chunker splits on `#`/`##`/`###` and a flat wall of text produces chunks with no
section path and therefore no meaningful citation.

Extraction is rejected rather than accepted empty. A scanned PDF with no text
layer, or a file whose content is images, yields almost nothing -- and an
"empty" document silently degrades every later answer. There is no OCR in
scope, so the honest outcome is a clear refusal at upload time.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.config import settings

#: Below this much extracted text a document cannot answer anything, so it is
#: refused with a reason instead of being indexed as filler.
MIN_EXTRACTED_CHARS = 200

MARKDOWN_SUFFIXES = {".md", ".markdown"}
TEXT_SUFFIXES = {".txt", ".text"}
HTML_SUFFIXES = {".html", ".htm"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}


class UnsupportedDocument(ValueError):
    """The file type is not one we can read."""


class EmptyDocument(ValueError):
    """The file was read but yielded too little text to be useful."""


@dataclass
class ParsedDocument:
    """The result of parsing one file."""

    markdown: str
    #: Suggested title: the first heading found, else the filename.
    title: str
    format: str
    headings: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.markdown)

    def preview(self, limit: int = 4000) -> str:
        """Truncated text for the /admin confirm-before-commit step."""
        if len(self.markdown) <= limit:
            return self.markdown
        return self.markdown[:limit].rsplit("\n", 1)[0] + "\n\n[... truncated ...]"

    def as_dict(self) -> dict:
        return {
            "title": self.title,
            "format": self.format,
            "char_count": self.char_count,
            "heading_count": len(self.headings),
            "headings": self.headings[:25],
            "warnings": self.warnings,
        }


def supported_suffixes() -> set[str]:
    return set(settings.allowed_upload_extensions)


def _tidy(markdown: str) -> str:
    markdown = markdown.replace("\r\n", "\n").replace("\r", "\n")
    markdown = re.sub(r"[ \t]+$", "", markdown, flags=re.MULTILINE)
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)
    return markdown.strip() + "\n"


def _headings(markdown: str) -> list[str]:
    return re.findall(r"^#{1,6}\s+(.+)$", markdown, flags=re.MULTILINE)


# ---------------------------------------------------------------------------
# Per-format parsers
# ---------------------------------------------------------------------------
def _parse_markdown(data: bytes) -> tuple[str, list[str]]:
    return data.decode("utf-8", errors="replace"), []


def _parse_text(data: bytes) -> tuple[str, list[str]]:
    """Plain text has no headings, so promote a leading title line to an H1.

    Without at least one heading the chunker cannot build a section path, and
    every chunk of the document would cite the same bare title.
    """
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    lines = text.split("\n")
    warnings = ["Plain text has no headings; the whole document is one section."]
    for index, line in enumerate(lines):
        if line.strip():
            if not line.startswith("#") and len(line.strip()) < 120:
                lines[index] = f"# {line.strip()}"
            break
    return "\n".join(lines), warnings


def _parse_html(data: bytes) -> tuple[str, list[str]]:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(data.decode("utf-8", errors="replace"), "html.parser")
    for element in soup.select("script, style, noscript, nav, header, footer, aside"):
        element.decompose()
    container = soup.find("main") or soup.find("article") or soup.body or soup
    markdown = markdownify(
        str(container),
        heading_style="ATX",
        strip=["a", "img", "sup"],
        escape_asterisks=False,
        escape_underscores=False,
        escape_misc=False,
    )
    return markdown, []


def _parse_pdf(data: bytes) -> tuple[str, list[str]]:
    """Extract a PDF's text layer, one `##` heading per page.

    Page headings are a pragmatic substitute for real structure: PDFs carry no
    reliable heading semantics, but without any headings the chunker cannot
    produce a section path. "Page 4" is a weak citation, yet it is honest and
    locatable, which beats citing the whole document.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    warnings: list[str] = []
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise EmptyDocument(f"The PDF is encrypted and cannot be read: {exc}")

    parts: list[str] = []
    empty_pages = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            warnings.append(f"Page {number} could not be read: {exc}")
            text = ""
        text = text.strip()
        if not text:
            empty_pages += 1
            continue
        parts.append(f"## Page {number}\n\n{text}")

    total = len(reader.pages)
    if empty_pages:
        warnings.append(
            f"{empty_pages} of {total} page(s) had no extractable text. "
            "Scanned or image-only PDFs need OCR, which is not supported."
        )
    return "\n\n".join(parts), warnings


def _parse_docx(data: bytes) -> tuple[str, list[str]]:
    """Map Word heading styles onto markdown headings.

    `Title` becomes `#` and `Heading N` becomes N+1 hashes, so `Heading 1`
    is `##` and `Heading 2` is `###`. Mapping both Title and Heading 1 to
    `#` would collapse two distinct levels together; shifting down by one
    mirrors how the curated documents are shaped (H1 = document title,
    H2 = section) and gives the chunker its full three usable levels.

    Tables are flattened to pipe rows.
    """
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    warnings: list[str] = []
    lines: list[str] = []
    styled_headings = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").strip() if paragraph.style else ""
        match = re.match(r"^Heading (\d)$", style)
        if match:
            level = min(int(match.group(1)) + 1, 6)
            lines.append(f"{'#' * level} {text}")
            styled_headings += 1
        elif style == "Title":
            lines.append(f"# {text}")
            styled_headings += 1
        elif style in {"List Bullet", "List Paragraph"}:
            lines.append(f"* {text}")
        elif style == "List Number":
            lines.append(f"1. {text}")
        else:
            lines.append(text)

    for table in document.tables:
        rows = [
            "| " + " | ".join(cell.text.strip() for cell in row.cells) + " |"
            for row in table.rows
        ]
        if rows:
            lines.append("")
            lines.extend(rows)

    if not styled_headings:
        warnings.append(
            "No Word heading styles found, so the document has no section "
            "structure. Applying Title / Heading 1 / Heading 2 styles before "
            "upload gives much more precise citations."
        )
        # Promote the first line to an H1 anyway: without a single heading
        # the chunker cannot build a section path, and every chunk would
        # cite the bare filename.
        for index, line in enumerate(lines):
            if line.strip() and not line.startswith(('#', '*', '1.', '|')):
                lines[index] = '# ' + line.strip()[:120]
                break
    return (chr(10) * 2).join(lines), warnings


_PARSERS = [
    (MARKDOWN_SUFFIXES, "markdown", _parse_markdown),
    (TEXT_SUFFIXES, "text", _parse_text),
    (HTML_SUFFIXES, "html", _parse_html),
    (PDF_SUFFIXES, "pdf", _parse_pdf),
    (DOCX_SUFFIXES, "docx", _parse_docx),
]


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def parse_to_markdown(data: bytes, filename: str) -> ParsedDocument:
    """Parse `data` according to `filename`'s extension.

    Raises `UnsupportedDocument` for an unknown extension and `EmptyDocument`
    when too little text comes out to be worth indexing.
    """
    suffix = Path(filename).suffix.lower()
    if not suffix:
        raise UnsupportedDocument(
            f"{filename!r} has no file extension, so its format is unknown. "
            f"Supported: {', '.join(sorted(supported_suffixes()))}"
        )

    for suffixes, format_name, parser in _PARSERS:
        if suffix not in suffixes:
            continue
        if suffix not in supported_suffixes():
            raise UnsupportedDocument(
                f"{suffix} files are not enabled. Supported: "
                f"{', '.join(sorted(supported_suffixes()))}"
            )
        try:
            raw, warnings = parser(data)
        except (EmptyDocument, UnsupportedDocument):
            raise
        except Exception as exc:
            raise EmptyDocument(
                f"Could not read the {format_name} file: "
                f"{type(exc).__name__}: {exc}"
            ) from exc

        markdown = _tidy(raw)
        if len(markdown.strip()) < MIN_EXTRACTED_CHARS:
            raise EmptyDocument(
                f"Only {len(markdown.strip())} characters of text could be "
                f"extracted from {filename!r} (at least {MIN_EXTRACTED_CHARS} "
                "are needed). If this is a scanned PDF it has no text layer, "
                "and OCR is not supported."
            )

        headings = _headings(markdown)
        title = headings[0] if headings else Path(filename).stem
        return ParsedDocument(
            markdown=markdown,
            title=title,
            format=format_name,
            headings=headings,
            warnings=warnings,
        )

    raise UnsupportedDocument(
        f"{suffix} files are not supported. Supported: "
        f"{', '.join(sorted(supported_suffixes()))}"
    )
